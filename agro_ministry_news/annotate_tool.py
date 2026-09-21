"""Local annotation tool for recording Orhan's own category labels.

Built 2026-09-21 at Orhan's request: raw CSV is not a workable editing surface for
500 rows of long Turkish article text (see agent_note_agroministrynews_FSOI.md,
Part 1, open decision #1 "Annotation UX").

Run it from agro_ministry_news/:

    python annotate_tool.py

Then open http://localhost:8000 in a browser if it doesn't open by itself.

WHICH FILES THIS TOUCHES - important:
  reads  agroforestministry_news_validation_sample.csv          (article text)
  reads  agroforestministry_news_validation_sample_CLAUDE_LABELS.csv   (never written)
  writes agroforestministry_news_validation_sample_ORHAN_LABELS.csv    (only this)

Orhan's labels live in their own separate file (his call, 2026-09-21: "I don't want
to change claude label file, let's keep it separate"). The CLAUDE_LABELS file is
opened read-only and is never modified, so nothing this tool does can damage the
classification work. Join the two on `Number` when you need them together.

Presentation order is RANDOM (fixed seed, so it's stable across restarts). The
sample's file order is batch order, and batches 1-2 were year-stratified while 3-5
were plain random - so file order is not uniformly random. Randomising means that if
annotation stops partway, whatever is done so far is still a representative subset of
the 500 rather than skewed toward the year-stratified early batches.

Nothing leaves this machine - the server binds to localhost only.
"""

import csv
import json
import os
import random
import shutil
import socket
import sys
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, HTTPServer

SAMPLE_CSV = "agroforestministry_news_validation_sample.csv"
CLAUDE_LABELS_CSV = "agroforestministry_news_validation_sample_CLAUDE_LABELS.csv"
ORHAN_LABELS_CSV = "agroforestministry_news_validation_sample_ORHAN_LABELS.csv"
ORHAN_FIELDNAMES = ["Number", "Orhan_Category", "Annotated_At"]

SHUFFLE_SEED = 42
PORT = 8000
PORT_TRIES = 10


class AnnotationStore:
    """Joins article text + Claude's labels for display, and owns Orhan's own labels file.

    The Claude labels file is read once and never written. The only file this class
    writes is ORHAN_LABELS_CSV.
    """

    def __init__(self, sample_csv=SAMPLE_CSV, claude_csv=CLAUDE_LABELS_CSV,
                 orhan_csv=ORHAN_LABELS_CSV):
        self.orhan_csv = orhan_csv
        self._backed_up = False
        self._stamps = {}

        with open(sample_csv, "r", encoding="utf-8-sig", newline="") as f:
            sample_rows = list(csv.DictReader(f))
        with open(claude_csv, "r", encoding="utf-8-sig", newline="") as f:
            claude_rows = list(csv.DictReader(f))

        claude_by_number = {r["Number"]: r for r in claude_rows}
        self.labels = self._load_orhan_labels(claude_rows)

        self.rows = []
        for art in sample_rows:
            num = art["Number"]
            lab = claude_by_number.get(num, {})
            self.rows.append({
                "Number": num,
                "Title": art.get("Title", ""),
                "Date": art.get("Date", ""),
                "URL": art.get("URL", ""),
                "Paragraphs": art.get("Paragraphs", ""),
                "Categories": lab.get("Categories", ""),
                "Ceremonial_Political": lab.get("Ceremonial_Political", ""),
                "Comment": lab.get("Comment", ""),
                "Orhan_Category": self.labels.get(num, ""),
            })

        # Stable random presentation order - same every run for a given row set.
        order = list(range(len(self.rows)))
        random.Random(SHUFFLE_SEED).shuffle(order)
        self.order = order
        self._by_number = {r["Number"]: r for r in self.rows}

    def _load_orhan_labels(self, claude_rows):
        """Reads the Orhan labels file. On first run, seeds it from any Orhan_Category
        values already sitting in the Claude labels file so nothing is stranded."""
        if os.path.isfile(self.orhan_csv):
            with open(self.orhan_csv, "r", encoding="utf-8-sig", newline="") as f:
                existing = list(csv.DictReader(f))
            # Keep the original annotation timestamps rather than restamping on rewrite.
            self._stamps = {r["Number"]: (r.get("Annotated_At") or "") for r in existing}
            return {r["Number"]: (r.get("Orhan_Category") or "") for r in existing}

        carried = {r["Number"]: (r.get("Orhan_Category") or "").strip()
                   for r in claude_rows if (r.get("Orhan_Category") or "").strip()}
        if carried:
            print(f"Carrying {len(carried)} existing Orhan_Category value(s) over from "
                  f"{CLAUDE_LABELS_CSV} into {self.orhan_csv} (that file is not modified).")
        return carried

    def ordered_rows(self):
        return [self.rows[i] for i in self.order]

    def done_count(self):
        return sum(1 for r in self.rows if r["Orhan_Category"].strip())

    def first_unannotated_position(self):
        """Position within the shuffled order, not the file order."""
        for pos, i in enumerate(self.order):
            if not self.rows[i]["Orhan_Category"].strip():
                return pos
        return 0

    def _backup_once(self):
        if self._backed_up or not os.path.isfile(self.orhan_csv):
            self._backed_up = True
            return
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(self.orhan_csv, f"{self.orhan_csv}.backup_{stamp}")
        self._backed_up = True

    def save(self, number, value):
        """Writes one label through to the Orhan labels file atomically."""
        number = str(number)
        row = self._by_number.get(number)
        if row is None:
            return False
        row["Orhan_Category"] = value
        if value.strip():
            self.labels[number] = value
        else:
            self.labels.pop(number, None)

        self._stamps[number] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self._backup_once()

        tmp = self.orhan_csv + ".tmp"
        with open(tmp, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=ORHAN_FIELDNAMES)
            writer.writeheader()
            # Written in sample file order so the file stays diff-friendly.
            for r in self.rows:
                num = r["Number"]
                v = self.labels.get(num, "")
                if v.strip():
                    writer.writerow({"Number": num, "Orhan_Category": v,
                                     "Annotated_At": self._stamps.get(num, "")})
        os.replace(tmp, self.orhan_csv)
        return True


PAGE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>FSOI annotation</title>
<style>
  :root {
    --bg:#f6f6f4; --panel:#fff; --ink:#1a1a18; --muted:#6b6b66;
    --line:#e0e0da; --accent:#4a6b4a; --accent-soft:#eef3ee; --warn:#8a6d3b;
  }
  @media (prefers-color-scheme: dark) {
    :root:not([data-theme="light"]) {
      --bg:#17171a; --panel:#1f1f23; --ink:#e8e8e4; --muted:#9a9a94;
      --line:#32323a; --accent:#8fb98f; --accent-soft:#242c24; --warn:#c9a86a;
    }
  }
  * { box-sizing:border-box; }
  body { margin:0; background:var(--bg); color:var(--ink);
    font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif; }
  header { position:sticky; top:0; z-index:10; background:var(--panel);
    border-bottom:1px solid var(--line); padding:10px 16px;
    display:flex; gap:14px; align-items:center; flex-wrap:wrap; }
  .grow { flex:1; }
  button { font:inherit; padding:7px 13px; border:1px solid var(--line);
    background:var(--panel); color:var(--ink); border-radius:7px; cursor:pointer; }
  button:hover:not(:disabled) { border-color:var(--accent); }
  button:disabled { opacity:.4; cursor:default; }
  .primary { background:var(--accent); border-color:var(--accent); color:#fff; }
  .bar { height:5px; background:var(--line); border-radius:3px; overflow:hidden; width:150px; }
  .bar i { display:block; height:100%; background:var(--accent); width:0; transition:width .2s; }
  .wrap { max-width:1500px; margin:0 auto; padding:16px;
    display:grid; grid-template-columns:1fr 430px; gap:16px; align-items:start; }
  @media (max-width:1000px) { .wrap { grid-template-columns:1fr; } }
  .card { background:var(--panel); border:1px solid var(--line);
    border-radius:10px; padding:16px; }
  .side { position:sticky; top:64px; }
  h1 { font-size:19px; line-height:1.35; margin:0 0 6px; }
  .meta { color:var(--muted); font-size:13px; margin-bottom:14px; }
  .meta a { color:var(--accent); }
  .body { white-space:pre-wrap; max-height:58vh; overflow-y:auto;
    padding-right:10px; font-size:14.5px; }
  .empty { color:var(--warn); font-style:italic; }
  h2 { font-size:11.5px; text-transform:uppercase; letter-spacing:.9px;
    color:var(--muted); margin:0 0 7px; font-weight:600; }
  .sect + .sect { margin-top:16px; padding-top:16px; border-top:1px solid var(--line); }
  .chip { display:inline-block; background:var(--accent-soft); color:var(--accent);
    border-radius:20px; padding:3px 10px; font-size:12.5px; margin:0 5px 5px 0; }
  .cmt { font-size:14px; color:var(--ink); }
  textarea { width:100%; min-height:130px; font:inherit; padding:11px;
    border:1px solid var(--line); border-radius:8px; resize:vertical;
    background:var(--bg); color:var(--ink); }
  textarea:focus { outline:2px solid var(--accent); outline-offset:-1px; }
  .status { font-size:12.5px; color:var(--muted); min-height:18px; margin-top:7px; }
  .saved { color:var(--accent); }
  .hint { font-size:12px; color:var(--muted); margin-top:10px; }
  kbd { background:var(--bg); border:1px solid var(--line); border-radius:4px;
    padding:1px 5px; font-size:11px; }
</style>
</head>
<body>
<header>
  <button id="prev">&larr; Prev</button>
  <button id="next">Next &rarr;</button>
  <strong id="pos"></strong>
  <div class="bar"><i id="fill"></i></div>
  <span id="done" style="color:var(--muted);font-size:13px"></span>
  <span class="grow"></span>
  <button id="skip">Next unannotated</button>
  <button id="xlsx" class="primary">Export .xlsx</button>
</header>

<div class="wrap">
  <div class="card">
    <h1 id="title"></h1>
    <div class="meta" id="meta"></div>
    <div class="body" id="text"></div>
  </div>

  <div>
    <div class="card side">
      <div class="sect">
        <h2>Claude's categories</h2>
        <div id="cats"></div>
      </div>
      <div class="sect">
        <h2>Ceremonial / political</h2>
        <div id="cer"></div>
      </div>
      <div class="sect">
        <h2>Claude's comment</h2>
        <div class="cmt" id="cmt"></div>
      </div>
      <div class="sect">
        <h2>Your category</h2>
        <textarea id="input" placeholder="Free text - not limited to the 44 codebook categories. This is where new candidate categories are meant to come from."></textarea>
        <div class="status" id="status"></div>
        <div class="hint">Saved to ORHAN_LABELS.csv - Claude's file is never touched.
          <kbd>Ctrl</kbd>+<kbd>&larr;</kbd>/<kbd>&rarr;</kbd> to move.</div>
      </div>
    </div>
  </div>
</div>

<script>
let rows = [], i = 0, timer = null;

function esc(s){ return (s||"").replace(/[&<>]/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;'}[c])); }

function render(){
  const r = rows[i];
  document.getElementById("pos").textContent = `${i+1} / ${rows.length}`;
  document.getElementById("title").textContent = r.Title || "(no title)";
  const link = r.URL ? ` &middot; <a href="${esc(r.URL)}" target="_blank" rel="noopener">open original</a>` : "";
  document.getElementById("meta").innerHTML =
    `Haber/${esc(r.Number)} &middot; ${esc(r.Date) || "no date"}${link}`;

  const t = document.getElementById("text");
  if ((r.Paragraphs||"").trim()) { t.textContent = r.Paragraphs; t.className = "body"; }
  else { t.textContent = "No article text was scraped for this row. This is usually a page the ministry removed - the site serves a 'content deleted' notice instead of a real article."; t.className = "body empty"; }

  const cats = (r.Categories||"").split(";").map(s=>s.trim()).filter(Boolean);
  document.getElementById("cats").innerHTML = cats.length
    ? cats.map(c=>`<span class="chip">${esc(c)}</span>`).join("")
    : `<span class="empty">none assigned</span>`;
  document.getElementById("cer").textContent = r.Ceremonial_Political || "-";
  document.getElementById("cmt").textContent = r.Comment || "-";

  document.getElementById("input").value = r.Orhan_Category || "";
  document.getElementById("status").textContent = "";
  document.getElementById("prev").disabled = i === 0;
  document.getElementById("next").disabled = i === rows.length - 1;
  progress();
}

function progress(){
  const n = rows.filter(r => (r.Orhan_Category||"").trim()).length;
  document.getElementById("fill").style.width = (100*n/rows.length) + "%";
  document.getElementById("done").textContent = `${n} annotated`;
}

function go(n){
  if (n < 0 || n >= rows.length) return;
  flush(); i = n; render();
  document.querySelector(".body").scrollTop = 0;
}

function flush(){
  if (timer) { clearTimeout(timer); timer = null; save(); }
}

function save(){
  const r = rows[i], v = document.getElementById("input").value;
  if (v === r.Orhan_Category) return;
  r.Orhan_Category = v;
  const st = document.getElementById("status");
  st.textContent = "saving...";
  fetch("/save", {
    method:"POST", headers:{"Content-Type":"application/json"},
    body: JSON.stringify({ number: r.Number, value: v })
  }).then(res => {
    if (!res.ok) throw new Error("save failed");
    st.textContent = "saved"; st.className = "status saved";
    progress();
  }).catch(() => {
    st.textContent = "SAVE FAILED - copy your text before navigating away";
    st.className = "status"; st.style.color = "var(--warn)";
  });
}

document.getElementById("input").addEventListener("input", () => {
  const st = document.getElementById("status");
  st.textContent = "typing..."; st.className = "status";
  clearTimeout(timer);
  timer = setTimeout(() => { timer = null; save(); }, 600);
});

document.getElementById("prev").onclick = () => go(i-1);
document.getElementById("next").onclick = () => go(i+1);
document.getElementById("skip").onclick = () => {
  flush();
  const n = rows.findIndex(r => !(r.Orhan_Category||"").trim());
  if (n === -1) alert("Every row has an annotation."); else go(n);
};
document.getElementById("xlsx").onclick = () => {
  flush();
  fetch("/export", {method:"POST"}).then(r=>r.json()).then(d=>alert("Written:\\n" + d.path));
};
document.addEventListener("keydown", e => {
  if (!e.ctrlKey) return;
  if (e.key === "ArrowRight") { e.preventDefault(); go(i+1); }
  if (e.key === "ArrowLeft")  { e.preventDefault(); go(i-1); }
});
window.addEventListener("beforeunload", flush);

fetch("/data").then(r=>r.json()).then(d => {
  rows = d.rows; i = d.start; render();
});
</script>
</body>
</html>
"""


def make_handler(store):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass  # keep the terminal readable; we print our own save lines

        def _send(self, code, body, ctype):
            data = body.encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", ctype + "; charset=utf-8")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self):
            if self.path == "/":
                self._send(200, PAGE, "text/html")
            elif self.path == "/data":
                payload = {"rows": store.ordered_rows(),
                           "start": store.first_unannotated_position()}
                self._send(200, json.dumps(payload), "application/json")
            else:
                self._send(404, "not found", "text/plain")

        def do_POST(self):
            if self.path == "/save":
                length = int(self.headers.get("Content-Length", 0))
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
                ok = store.save(payload["number"], payload["value"])
                if ok:
                    print(f"  [{store.done_count()}/{len(store.rows)}] "
                          f"Haber/{payload['number']} -> {store.orhan_csv}")
                self._send(200 if ok else 400, json.dumps({"ok": ok}), "application/json")
            elif self.path == "/export":
                path = export_xlsx(store)
                self._send(200, json.dumps({"path": os.path.abspath(path)}),
                           "application/json")
            else:
                self._send(404, "not found", "text/plain")

    return Handler


def export_xlsx(store, path="agroforestministry_news_validation_sample_ANNOTATED.xlsx"):
    """Convenience export joining article + Claude's labels + Orhan's labels for reading.
    A formatted snapshot, never read back in - the CSVs remain the source of truth."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "annotations"
    headers = ["Number", "Date", "Title", "Categories",
               "Ceremonial_Political", "Comment", "Orhan_Category"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in store.rows:
        ws.append([r.get(h, "") for h in headers])
    for col, width in zip("ABCDEFG", [10, 12, 55, 34, 10, 60, 34]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def find_free_port(start=PORT, tries=PORT_TRIES):
    """Returns the first free port at or after `start`.

    Re-running the tool while an earlier copy is still alive is the normal case
    (closing the browser tab does not stop the server), so we move to the next port
    instead of failing with a confusing 'address in use' traceback.
    """
    for p in range(start, start + tries):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", p)) != 0:
                return p
    return None


def main():
    # Work relative to this file, so it behaves the same whether it is launched from
    # a terminal, from VS Code's Run button, or by double-clicking.
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    for f in (SAMPLE_CSV, CLAUDE_LABELS_CSV):
        if not os.path.isfile(f):
            raise SystemExit(f"Missing {f} in {os.getcwd()}")

    store = AnnotationStore()
    port = find_free_port()
    if port is None:
        raise SystemExit(
            f"Ports {PORT}-{PORT + PORT_TRIES - 1} are all busy. An older copy of this "
            f"tool is probably still running - close those terminals and try again."
        )

    url = f"http://localhost:{port}"
    print(f"Loaded {len(store.rows)} articles ({store.done_count()} already annotated).")
    print(f"Your labels    -> {ORHAN_LABELS_CSV}")
    print(f"Claude's labels-> {CLAUDE_LABELS_CSV}  (read-only, never modified)")
    print("Order          -> random, fixed seed (same every run)")
    if port != PORT:
        print(f"\nPort {PORT} was busy, using {port} instead.")
    print(f"\n    {url}\n")
    print("If the browser does not open by itself, copy that address into one.")
    print("Press Ctrl+C here to stop the server.\n")

    server = HTTPServer(("127.0.0.1", port), make_handler(store))
    try:
        webbrowser.open(url)
    except Exception:
        pass  # not fatal - the URL is printed above
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print(f"\nStopped. {store.done_count()}/{len(store.rows)} annotated.")
        sys.exit(0)


if __name__ == "__main__":
    main()
